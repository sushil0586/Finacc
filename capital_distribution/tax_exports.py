from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


LINE_HEADERS = (
    "Line ID",
    "Source Line",
    "Stakeholder",
    "Component",
    "Treatment",
    "Book Amount",
    "Calculated Allowable",
    "Calculated Disallowed",
    "Override Allowable",
    "Final Allowable",
    "Final Disallowed",
    "Override Reason",
    "Evidence References",
)
MONEY_COLUMNS = {6, 7, 8, 9, 10, 11}


def _safe_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, sort_keys=True, separators=(",", ":"), default=str)
    text = str(value)
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{text}"
    return text


def _iso(value) -> str:
    return value.isoformat() if value else ""


def build_tax_working_export(working) -> dict:
    policy = working.policy_snapshot or {}
    source = working.source_snapshot or {}
    metadata = [
        ("Working ID", working.id),
        ("Entity", working.entity.entityname),
        ("Entity ID", working.entity_id),
        ("Financial Year", working.entityfin.desc or working.entityfin_id),
        ("Financial Year ID", working.entityfin_id),
        ("Branch", working.subentity.subentityname if working.subentity_id else "All branches"),
        ("Branch ID", working.subentity_id or ""),
        ("Period From", working.period_from.isoformat()),
        ("Period To", working.period_to.isoformat()),
        ("Status", working.status),
        ("Source Run", working.run_id),
        ("Source Run Hash", source.get("run_calculation_hash", "")),
        ("Source Posting Batch", source.get("posting_batch", "")),
        ("Tax Policy", policy.get("policy_code", working.tax_policy.policy_code)),
        ("Tax Policy Version", policy.get("version_number", working.tax_policy.version_number)),
        ("Statutory Reference", policy.get("statutory_reference", "")),
        ("Statutory Source", policy.get("source_url", "")),
        ("Calculation Hash", working.calculation_hash),
        ("Calculated At", _iso(working.calculated_at)),
        ("Calculated By", working.calculated_by_id or ""),
        ("Submitted At", _iso(working.submitted_at)),
        ("Submitted By", working.submitted_by_id or ""),
        ("Approved At", _iso(working.approved_at)),
        ("Approved By", working.approved_by_id or ""),
        ("Reversed At", _iso(working.reversed_at)),
        ("Reversed By", working.reversed_by_id or ""),
        ("Lifecycle Reason", working.lifecycle_reason),
    ]
    totals = [
        ("Book Amount", working.book_amount),
        ("Allowable Amount", working.allowable_amount),
        ("Disallowed Amount", working.disallowed_amount),
    ]
    rows = [
        (
            line.id,
            line.source_line_id,
            line.stakeholder_name,
            line.component_type,
            line.treatment,
            line.book_amount,
            line.calculated_allowable_amount,
            line.calculated_disallowed_amount,
            line.override_allowable_amount,
            line.allowable_amount,
            line.disallowed_amount,
            line.override_reason,
            line.evidence_references,
        )
        for line in working.lines.all()
    ]
    return {"metadata": metadata, "totals": totals, "headers": LINE_HEADERS, "rows": rows}


def tax_working_filename(working, extension: str) -> str:
    return f"tax_working_{working.id}_{working.period_from:%Y%m%d}_{working.period_to:%Y%m%d}.{extension}"


def render_tax_working_csv(payload: dict) -> bytes:
    stream = StringIO()
    writer = csv.writer(stream)
    writer.writerow(["Tax Working Summary"])
    writer.writerows((_safe_text(label), _safe_text(value)) for label, value in payload["metadata"])
    writer.writerow([])
    writer.writerow(["Reconciliation Totals"])
    writer.writerows((_safe_text(label), _safe_text(value)) for label, value in payload["totals"])
    writer.writerow([])
    writer.writerow(payload["headers"])
    writer.writerows([_safe_text(value) for value in row] for row in payload["rows"])
    return stream.getvalue().encode("utf-8-sig")


def _style_header(row):
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def render_tax_working_xlsx(payload: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Tax Working Summary"])
    _style_header(summary[1])
    for label, value in payload["metadata"]:
        summary.append([_safe_text(label), _safe_text(value)])
    summary.append([])
    summary.append(["Reconciliation Totals"])
    _style_header(summary[summary.max_row])
    for label, value in payload["totals"]:
        summary.append([_safe_text(label), float(value)])
        summary.cell(summary.max_row, 2).number_format = "#,##0.00"
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 72
    summary.freeze_panes = "A2"

    lines = workbook.create_sheet("Lines")
    lines.append(list(payload["headers"]))
    _style_header(lines[1])
    for source_row in payload["rows"]:
        row = []
        for index, value in enumerate(source_row, start=1):
            if index in MONEY_COLUMNS and value is not None:
                row.append(float(value))
            else:
                row.append(_safe_text(value))
        lines.append(row)
    for column in MONEY_COLUMNS:
        for cell in lines.iter_cols(min_col=column, max_col=column, min_row=2):
            for value in cell:
                value.number_format = "#,##0.00"
    for index, header in enumerate(payload["headers"], start=1):
        lines.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 3, 14), 42)
    lines.auto_filter.ref = lines.dimensions
    lines.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_tax_working_pdf(payload: dict) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title="Tax Working",
    )
    styles = getSampleStyleSheet()
    small = styles["BodyText"]
    small.fontSize = 7
    small.leading = 9
    metadata_rows = [[Paragraph(f"<b>{escape(_safe_text(label))}</b>", small), Paragraph(escape(_safe_text(value)), small)] for label, value in payload["metadata"]]
    total_rows = [[Paragraph(f"<b>{escape(_safe_text(label))}</b>", small), f"{value:,.2f}"] for label, value in payload["totals"]]
    line_rows = [[Paragraph(escape(_safe_text(value)), small) for value in row] for row in payload["rows"]]
    story = [Paragraph("Tax Working", styles["Title"]), Spacer(1, 4)]
    metadata_table = Table(metadata_rows, colWidths=[48 * mm, 210 * mm])
    metadata_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF0F8"))]))
    story.extend([metadata_table, Spacer(1, 8), Paragraph("Reconciliation Totals", styles["Heading2"])])
    totals_table = Table(total_rows, colWidths=[48 * mm, 45 * mm])
    totals_table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("ALIGN", (1, 0), (1, -1), "RIGHT")]))
    story.extend([totals_table, PageBreak(), Paragraph("Tax Working Lines", styles["Heading2"])])
    widths = [12, 15, 24, 20, 16, 18, 21, 21, 20, 18, 18, 28, 37]
    table = Table([[Paragraph(f"<b>{header}</b>", small) for header in payload["headers"]]] + line_rows, colWidths=[width * mm for width in widths], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F5597")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")])]))
    story.append(table)
    document.build(story)
    return buffer.getvalue()
