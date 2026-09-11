from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from rest_framework.exceptions import ValidationError as DRFValidationError
from rest_framework.test import APIClient, APITestCase

from Authentication.models import User
from entity.models import (
    Entity,
    EntityFinancialYear,
    EntityOwnershipV2,
    EntityTaxProfile,
    GstRegistrationType,
    SubEntity,
)
from financial.models import Ledger, account
from posting.models import EntityStaticAccountMap, JournalLine, StaticAccount, StaticAccountGroup, TxnType
from reports.services.controls.opening_generation import _build_opening_lines
from reports.services.controls.year_end_close import _build_close_journal_lines, _posted_appropriation_coverage

from .models import (
    CapitalDistributionAuditEvent,
    CapitalDistributionAccountMapping,
    CapitalDistributionLine,
    CapitalDistributionRun,
    DistributionPolicyVersion,
    EntityFormationProfile,
    FormationType,
)
from .calculations import calculate_segment
from .reporting import build_appropriation_statement
from .services import (
    approve_policy,
    create_policy,
    materialize_formation_profile,
    resolve_entity_formation,
    seed_policy_from_ownership,
    submit_policy,
    calculate_distribution_run,
    approve_distribution_run,
    post_distribution_run,
    reverse_distribution_run,
    submit_distribution_run,
    upsert_account_mapping,
)


def aware(year, month, day):
    return timezone.make_aware(datetime(year, month, day))


class CapitalDistributionFixtureMixin:
    def make_user(self, suffix):
        return User.objects.create_user(
            username=f"capital-{suffix}",
            email=f"capital-{suffix}@example.com",
            password="secret123",
        )

    def make_entity(self, *, name="Capital Entity", owner=None, business_type=Entity.BusinessType.MIXED):
        owner = owner or self.make_user(name.lower().replace(" ", "-"))
        gst_type, _ = GstRegistrationType.objects.get_or_create(Name="Regular", defaults={"Description": "Regular"})
        return Entity.objects.create(
            entityname=name,
            entitydesc=name,
            legalname=name,
            GstRegitrationType=gst_type,
            business_type=business_type,
            createdby=owner,
        )

    def add_owner(self, entity, *, name, ownership_type, share):
        return EntityOwnershipV2.objects.create(
            entity=entity,
            name=name,
            ownership_type=ownership_type,
            share_percentage=share,
            account_preference=EntityOwnershipV2.AccountPreference.CURRENT,
            createdby=entity.createdby,
        )

    def make_partnership(self, *, name="Partnership Entity", owner=None):
        entity = self.make_entity(name=name, owner=owner)
        first = self.add_owner(
            entity,
            name="Partner A",
            ownership_type=EntityOwnershipV2.OwnershipType.PARTNER,
            share="60.0000",
        )
        second = self.add_owner(
            entity,
            name="Partner B",
            ownership_type=EntityOwnershipV2.OwnershipType.PARTNER,
            share="40.0000",
        )
        return entity, first, second

    def policy_payload(self, first, second, **overrides):
        payload = {
            "effective_from": date(2026, 4, 1),
            "effective_to": date(2027, 3, 31),
            "governing_document_reference": "DEED-2026-01",
            "configuration": {"run_cadence": "annual", "rounding": "half_up"},
            "stakeholders": [
                {
                    "ownership": first,
                    "target_type": "partner",
                    "target_name": first.name,
                    "profit_percentage": "60.0000",
                    "loss_percentage": "60.0000",
                },
                {
                    "ownership": second,
                    "target_type": "partner",
                    "target_name": second.name,
                    "profit_percentage": "40.0000",
                    "loss_percentage": "40.0000",
                },
            ],
        }
        payload.update(overrides)
        return payload


class FormationResolutionTests(CapitalDistributionFixtureMixin, TestCase):
    def test_partnership_is_supported_from_consistent_ownership(self):
        entity, _, _ = self.make_partnership()

        result = resolve_entity_formation(entity)

        self.assertEqual(result.formation_type, FormationType.PARTNERSHIP)
        self.assertEqual(result.status, EntityFormationProfile.Status.VERIFIED)
        self.assertTrue(result.supported)
        self.assertEqual(result.strategy_code, "partnership_appropriation_v1")

    def test_llpin_outranks_compatible_partner_rows(self):
        entity, _, _ = self.make_partnership(name="LLP Entity")
        EntityTaxProfile.objects.create(entity=entity, llpin_no="AAA-1234", createdby=entity.createdby)

        result = resolve_entity_formation(entity)

        self.assertEqual(result.formation_type, FormationType.LLP)
        self.assertEqual(result.status, EntityFormationProfile.Status.VERIFIED)
        self.assertTrue(result.supported)

    def test_cin_with_partner_rows_is_contradictory(self):
        entity, _, _ = self.make_partnership(name="Contradictory Entity")
        EntityTaxProfile.objects.create(
            entity=entity,
            cin_no="U12345PB2026PTC123456",
            createdby=entity.createdby,
        )

        result = resolve_entity_formation(entity)

        self.assertEqual(result.formation_type, FormationType.COMPANY)
        self.assertEqual(result.status, EntityFormationProfile.Status.CONTRADICTORY)
        self.assertFalse(result.supported)
        self.assertIn("formation_ownership_conflict", {row["code"] for row in result.readiness_issues})

    def test_recognized_company_is_gated_until_later_wave(self):
        entity = self.make_entity(name="Company Entity")
        EntityTaxProfile.objects.create(
            entity=entity,
            cin_no="U12345PB2026PTC123456",
            createdby=entity.createdby,
        )
        self.add_owner(
            entity,
            name="Shareholder A",
            ownership_type=EntityOwnershipV2.OwnershipType.SHAREHOLDER,
            share="100.0000",
        )

        result = resolve_entity_formation(entity)

        self.assertEqual(result.formation_type, FormationType.COMPANY)
        self.assertEqual(result.status, EntityFormationProfile.Status.UNSUPPORTED)
        self.assertFalse(result.supported)
        self.assertIn("strategy_not_enabled", {row["code"] for row in result.readiness_issues})

    def test_business_type_only_requires_verification_and_is_not_wave_one(self):
        entity = self.make_entity(name="NGO Entity", business_type=Entity.BusinessType.NGO)

        result = resolve_entity_formation(entity)

        self.assertEqual(result.formation_type, FormationType.NGO)
        self.assertEqual(result.status, EntityFormationProfile.Status.UNSUPPORTED)
        self.assertFalse(result.supported)
        self.assertIn("business_type_only", {row["code"] for row in result.readiness_issues})

    def test_materialization_is_idempotent_until_evidence_changes(self):
        entity, _, _ = self.make_partnership()

        first = materialize_formation_profile(entity=entity, actor=entity.createdby)
        repeated = materialize_formation_profile(entity=entity, actor=entity.createdby)
        self.add_owner(
            entity,
            name="Director Conflict",
            ownership_type=EntityOwnershipV2.OwnershipType.DIRECTOR,
            share="0.0000",
        )
        changed = materialize_formation_profile(entity=entity, actor=entity.createdby)

        self.assertEqual(first.id, repeated.id)
        self.assertEqual(changed.version_number, 2)
        self.assertEqual(changed.status, EntityFormationProfile.Status.CONTRADICTORY)
        self.assertEqual(CapitalDistributionAuditEvent.objects.filter(action="formation_resolved").count(), 2)


class DistributionPolicyServiceTests(CapitalDistributionFixtureMixin, TestCase):
    def setUp(self):
        self.maker = self.make_user("maker")
        self.approver = self.make_user("approver")
        self.entity, self.first, self.second = self.make_partnership(owner=self.maker)
        self.profile = materialize_formation_profile(entity=self.entity, actor=self.maker)

    def test_policy_lifecycle_is_versioned_audited_and_maker_checker(self):
        policy = create_policy(
            entity=self.entity,
            formation_profile=self.profile,
            payload=self.policy_payload(self.first, self.second),
            actor=self.maker,
        )
        policy = submit_policy(policy=policy, actor=self.maker, reason="Ready for review")

        with self.assertRaises(ValidationError):
            approve_policy(policy=policy, actor=self.maker, reason="Self approve")

        policy = approve_policy(policy=policy, actor=self.approver, reason="Checked against deed")

        self.assertEqual(policy.status, DistributionPolicyVersion.Status.APPROVED)
        self.assertEqual(policy.submitted_by, self.maker)
        self.assertEqual(policy.approved_by, self.approver)
        self.assertEqual(policy.stakeholders.count(), 2)
        self.assertEqual(
            list(policy.audit_events.values_list("action", flat=True).order_by("created_at", "id")),
            ["policy_created", "policy_submitted", "policy_approved"],
        )

    def test_invalid_share_total_rolls_back_policy_and_rows(self):
        payload = self.policy_payload(self.first, self.second)
        payload["stakeholders"][1]["profit_percentage"] = "30.0000"

        with self.assertRaises(ValidationError):
            create_policy(
                entity=self.entity,
                formation_profile=self.profile,
                payload=payload,
                actor=self.maker,
            )

        self.assertFalse(DistributionPolicyVersion.objects.exists())
        self.assertFalse(CapitalDistributionAuditEvent.objects.filter(action="policy_created").exists())

    def test_approved_policy_period_cannot_overlap(self):
        first_policy = create_policy(
            entity=self.entity,
            formation_profile=self.profile,
            payload=self.policy_payload(self.first, self.second),
            actor=self.maker,
        )
        submit_policy(policy=first_policy, actor=self.maker)
        approve_policy(policy=first_policy, actor=self.approver)

        with self.assertRaises(ValidationError):
            create_policy(
                entity=self.entity,
                formation_profile=self.profile,
                payload=self.policy_payload(
                    self.first,
                    self.second,
                    effective_from=date(2026, 10, 1),
                    effective_to=date(2027, 9, 30),
                ),
                actor=self.maker,
            )

        self.assertEqual(DistributionPolicyVersion.objects.count(), 1)

    def test_foreign_ownership_row_is_rejected_atomically(self):
        other_entity = self.make_entity(name="Other Entity")
        foreign_owner = self.add_owner(
            other_entity,
            name="Foreign Partner",
            ownership_type=EntityOwnershipV2.OwnershipType.PARTNER,
            share="40.0000",
        )
        payload = self.policy_payload(self.first, foreign_owner)

        with self.assertRaises(ValidationError):
            create_policy(
                entity=self.entity,
                formation_profile=self.profile,
                payload=payload,
                actor=self.maker,
            )

        self.assertFalse(DistributionPolicyVersion.objects.exists())

    def test_seed_policy_uses_active_ownership_ratios(self):
        retired = self.add_owner(
            self.entity,
            name="Retired Partner",
            ownership_type=EntityOwnershipV2.OwnershipType.PARTNER,
            share="25.0000",
        )
        retired.effective_to = date(2026, 3, 31)
        retired.save(update_fields=("effective_to", "updated_at"))

        policy = seed_policy_from_ownership(
            entity=self.entity,
            formation_profile=self.profile,
            payload={
                "effective_from": date(2026, 4, 1),
                "effective_to": date(2027, 3, 31),
                "governing_document_reference": "DEED-SEED-01",
            },
            actor=self.maker,
        )

        self.assertEqual(policy.stakeholders.count(), 2)
        self.assertEqual(
            list(policy.stakeholders.values_list("profit_percentage", flat=True)),
            [Decimal("60.0000"), Decimal("40.0000")],
        )
        self.assertTrue(policy.audit_events.filter(action="policy_seeded_from_ownership").exists())

    def test_policy_rejects_partner_retiring_inside_policy_period(self):
        self.second.effective_to = date(2026, 9, 30)
        self.second.save(update_fields=("effective_to", "updated_at"))

        with self.assertRaises(ValidationError):
            create_policy(
                entity=self.entity,
                formation_profile=self.profile,
                payload=self.policy_payload(self.first, self.second),
                actor=self.maker,
            )

        self.assertFalse(DistributionPolicyVersion.objects.exists())


@override_settings(ROOT_URLCONF="FA.urls", AUTH_PASSWORD_VALIDATORS=[])
class CapitalDistributionAPITests(CapitalDistributionFixtureMixin, APITestCase):
    def setUp(self):
        self.maker = self.make_user("api-maker")
        self.approver = self.make_user("api-approver")
        self.entity, self.first, self.second = self.make_partnership(name="API Partnership", owner=self.maker)
        self.other_entity, _, _ = self.make_partnership(name="Other API Partnership", owner=self.maker)
        self.entityfin = EntityFinancialYear.objects.create(
            entity=self.entity,
            desc="FY 2026-27",
            year_code="FY2026-27",
            finstartyear=aware(2026, 4, 1),
            finendyear=aware(2027, 3, 31),
            createdby=self.maker,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.maker)
        self.subscription_patch = patch(
            "capital_distribution.views.SubscriptionService.assert_entity_access",
            return_value=None,
        )
        self.permission_patch = patch(
            "capital_distribution.views.EffectivePermissionService.permission_codes_for_user",
            return_value={
                "capital_distribution.setup.view",
                "capital_distribution.setup.manage",
                "capital_distribution.policy.view",
                "capital_distribution.policy.manage",
                "capital_distribution.policy.submit",
                "capital_distribution.policy.approve",
                "capital_distribution.run.view",
                "capital_distribution.run.calculate",
            },
        )
        self.subscription_patch.start()
        self.permission_mock = self.permission_patch.start()
        self.addCleanup(self.subscription_patch.stop)
        self.addCleanup(self.permission_patch.stop)

    def test_resolve_create_list_and_cross_entity_detail_isolation(self):
        formation_response = self.client.post(
            reverse("capital_distribution_api:formation-profile"),
            {"entity": self.entity.id, "effective_from": "2026-04-01"},
            format="json",
        )
        self.assertEqual(formation_response.status_code, 201, formation_response.data)
        profile_id = formation_response.data["id"]

        create_response = self.client.post(
            reverse("capital_distribution_api:policy-list"),
            {
                "entity": self.entity.id,
                "entityfinid": self.entityfin.id,
                "formation_profile": profile_id,
                "effective_from": "2026-04-01",
                "effective_to": "2027-03-31",
                "governing_document_reference": "DEED-API-01",
                "configuration": {"run_cadence": "annual"},
                "stakeholders": [
                    {
                        "ownership": self.first.id,
                        "target_type": "partner",
                        "target_name": self.first.name,
                        "profit_percentage": "60.0000",
                    },
                    {
                        "ownership": self.second.id,
                        "target_type": "partner",
                        "target_name": self.second.name,
                        "profit_percentage": "40.0000",
                    },
                ],
            },
            format="json",
        )
        self.assertEqual(create_response.status_code, 201, create_response.data)
        policy_id = create_response.data["id"]

        list_response = self.client.get(
            reverse("capital_distribution_api:policy-list"),
            {"entity": self.entity.id, "entityfinid": self.entityfin.id},
        )
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(list_response.data["count"], 1)

        submit_response = self.client.post(
            reverse("capital_distribution_api:policy-submit", kwargs={"policy_id": policy_id}),
            {
                "entity": self.entity.id,
                "expected_updated_at": create_response.data["updated_at"],
                "reason": "API maker review complete",
            },
            format="json",
        )
        self.assertEqual(submit_response.status_code, 200, submit_response.data)
        self.assertEqual(submit_response.data["status"], DistributionPolicyVersion.Status.SUBMITTED)

        self_approve_response = self.client.post(
            reverse("capital_distribution_api:policy-approve", kwargs={"policy_id": policy_id}),
            {
                "entity": self.entity.id,
                "expected_updated_at": submit_response.data["updated_at"],
                "reason": "Must not self approve",
            },
            format="json",
        )
        self.assertEqual(self_approve_response.status_code, 400)

        self.client.force_authenticate(self.approver)
        approve_response = self.client.post(
            reverse("capital_distribution_api:policy-approve", kwargs={"policy_id": policy_id}),
            {
                "entity": self.entity.id,
                "expected_updated_at": submit_response.data["updated_at"],
                "reason": "Independent API approval",
            },
            format="json",
        )
        self.assertEqual(approve_response.status_code, 200, approve_response.data)
        self.assertEqual(approve_response.data["status"], DistributionPolicyVersion.Status.APPROVED)

        foreign_response = self.client.get(
            reverse("capital_distribution_api:policy-detail", kwargs={"policy_id": policy_id}),
            {"entity": self.other_entity.id},
        )
        self.assertEqual(foreign_response.status_code, 400)
        self.assertIn("policy", foreign_response.data)

    def test_stale_patch_returns_conflict_without_mutation(self):
        profile = materialize_formation_profile(entity=self.entity, actor=self.maker)
        policy = create_policy(
            entity=self.entity,
            formation_profile=profile,
            payload=self.policy_payload(
                self.first,
                self.second,
                entityfin=self.entityfin,
            ),
            actor=self.maker,
        )
        stale_version = policy.updated_at
        policy.notes = "Updated elsewhere"
        policy.save(update_fields=("notes", "updated_at"))

        response = self.client.patch(
            reverse("capital_distribution_api:policy-detail", kwargs={"policy_id": policy.id}),
            {
                "entity": self.entity.id,
                "expected_updated_at": stale_version.isoformat(),
                "notes": "Stale overwrite",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 409, response.data)
        policy.refresh_from_db()
        self.assertEqual(policy.notes, "Updated elsewhere")

    def test_compare_reports_business_changes_and_enforces_entity_scope(self):
        profile = materialize_formation_profile(entity=self.entity, actor=self.maker)
        first = create_policy(
            entity=self.entity,
            formation_profile=profile,
            payload=self.policy_payload(self.first, self.second, notes="Original"),
            actor=self.maker,
        )
        second = create_policy(
            entity=self.entity,
            formation_profile=profile,
            payload=self.policy_payload(self.first, self.second, notes="Amended"),
            actor=self.maker,
        )

        response = self.client.get(
            reverse(
                "capital_distribution_api:policy-compare",
                kwargs={"policy_id": first.id, "other_policy_id": second.id},
            ),
            {"entity": self.entity.id},
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(set(response.data["changes"]), {"notes"})

        foreign_response = self.client.get(
            reverse(
                "capital_distribution_api:policy-compare",
                kwargs={"policy_id": first.id, "other_policy_id": second.id},
            ),
            {"entity": self.other_entity.id},
        )
        self.assertEqual(foreign_response.status_code, 400)

    def test_api_denies_missing_permission(self):
        self.permission_mock.return_value = set()

        response = self.client.get(
            reverse("capital_distribution_api:formation-profile"),
            {"entity": self.entity.id},
        )

        self.assertEqual(response.status_code, 403)

    def test_branch_view_can_read_shared_setup_only_with_an_allowed_branch_context(self):
        branch = SubEntity.objects.create(entity=self.entity, subentityname="Branch A")
        profile = materialize_formation_profile(entity=self.entity, actor=self.maker)
        policy = create_policy(
            entity=self.entity,
            formation_profile=profile,
            payload=self.policy_payload(self.first, self.second, entityfin=self.entityfin),
            actor=self.maker,
        )

        with patch(
            "capital_distribution.views.EffectivePermissionService.has_scope_access",
            side_effect=lambda _user, _entity_id, subentity_id=None: subentity_id == branch.id,
        ):
            formation = self.client.get(
                reverse("capital_distribution_api:formation-profile"),
                {"entity": self.entity.id, "entityfinid": self.entityfin.id, "subentity": branch.id},
            )
            self.assertEqual(formation.status_code, 200, formation.data)

            policies = self.client.get(
                reverse("capital_distribution_api:policy-list"),
                {"entity": self.entity.id, "entityfinid": self.entityfin.id, "subentity": branch.id},
            )
            self.assertEqual(policies.status_code, 200, policies.data)
            self.assertEqual([row["id"] for row in policies.data["results"]], [policy.id])

            detail = self.client.get(
                reverse("capital_distribution_api:policy-detail", kwargs={"policy_id": policy.id}),
                {"entity": self.entity.id, "subentity": branch.id},
            )
            self.assertEqual(detail.status_code, 200, detail.data)

            unscoped = self.client.get(
                reverse("capital_distribution_api:formation-profile"),
                {"entity": self.entity.id, "entityfinid": self.entityfin.id},
            )
            self.assertEqual(unscoped.status_code, 403, unscoped.data)


class PartnershipCalculationTests(CapitalDistributionFixtureMixin, TestCase):
    def setUp(self):
        self.maker = self.make_user("calc-maker")
        self.approver = self.make_user("calc-approver")
        self.entity, self.first, self.second = self.make_partnership(name="Calculation Partnership", owner=self.maker)
        self.entityfin = EntityFinancialYear.objects.create(
            entity=self.entity,
            desc="FY 2026-27",
            year_code="FY2026-27",
            finstartyear=aware(2026, 4, 1),
            finendyear=aware(2027, 3, 31),
            createdby=self.maker,
        )
        self.profile = materialize_formation_profile(entity=self.entity, actor=self.maker)
        self.policy = create_policy(
            entity=self.entity,
            formation_profile=self.profile,
            payload=self.policy_payload(self.first, self.second, entityfin=self.entityfin),
            actor=self.maker,
        )
        first_row, second_row = self.policy.stakeholders.order_by("sort_order")
        first_row.configuration = {
            "remuneration": {"enabled": True, "method": "fixed", "amount": "12000", "prorate": True},
            "capital_interest": {"enabled": True, "rate": "10"},
            "drawing_interest": {"enabled": True, "rate": "12"},
        }
        first_row.save(update_fields=("configuration", "updated_at"))
        self.first_row = first_row
        self.second_row = second_row

    def approve(self):
        submit_policy(policy=self.policy, actor=self.maker)
        self.policy = approve_policy(policy=self.policy, actor=self.approver)

    def test_golden_annual_calculation_matches_expected_components(self):
        result = calculate_segment(
            source_profit=Decimal("100000"),
            book_adjustments=Decimal("0"),
            period_from=date(2026, 4, 1),
            period_to=date(2027, 3, 31),
            stakeholder_rows=[self.first_row, self.second_row],
            balances={
                self.first_row.id: {"capital_balance": "100000", "drawing_balance": "10000"},
                self.second_row.id: {"capital_balance": "0", "drawing_balance": "0"},
            },
        )
        amounts = {(line.stakeholder_id, line.component_type): line.amount for line in result.lines}
        self.assertEqual(amounts[(self.first_row.id, "remuneration")], Decimal("12000.00"))
        self.assertEqual(amounts[(self.first_row.id, "capital_interest")], Decimal("10000.00"))
        self.assertEqual(amounts[(self.first_row.id, "drawing_interest")], Decimal("1200.00"))
        self.assertEqual(result.distributable_result, Decimal("79200.00"))
        self.assertEqual(amounts[(self.first_row.id, "residual_profit")], Decimal("47520.00"))
        self.assertEqual(amounts[(self.second_row.id, "residual_profit")], Decimal("31680.00"))

    def test_loss_uses_loss_ratio_and_debits_stakeholders(self):
        self.first_row.loss_percentage = Decimal("25")
        self.second_row.loss_percentage = Decimal("75")
        self.first_row.configuration = {}
        result = calculate_segment(
            source_profit=Decimal("-100.01"),
            book_adjustments=Decimal("0"),
            period_from=date(2026, 4, 1),
            period_to=date(2026, 4, 30),
            stakeholder_rows=[self.first_row, self.second_row],
            balances={},
        )
        residual = [line for line in result.lines if line.component_type == "residual_loss"]
        self.assertEqual([line.amount for line in residual], [Decimal("25.00"), Decimal("75.01")])
        self.assertTrue(all(line.side == CapitalDistributionLine.Side.DEBIT for line in residual))

    def test_actual_actual_uses_366_days_in_leap_year(self):
        self.first_row.configuration = {"capital_interest": {"enabled": True, "rate": "10"}}
        result = calculate_segment(
            source_profit=Decimal("0"),
            book_adjustments=Decimal("0"),
            period_from=date(2024, 1, 1),
            period_to=date(2024, 1, 31),
            stakeholder_rows=[self.first_row, self.second_row],
            balances={self.first_row.id: {"capital_balance": "100000"}},
            day_count_convention="actual_actual",
        )
        interest = next(line for line in result.lines if line.component_type == "capital_interest")
        self.assertEqual(interest.amount, Decimal("846.99"))

    def test_percentage_remuneration_uses_period_profit_without_double_proration(self):
        self.first_row.configuration = {
            "remuneration": {
                "enabled": True,
                "method": "percentage_of_book_profit",
                "rate": "10",
                "prorate": True,
            }
        }
        result = calculate_segment(
            source_profit=Decimal("10000"),
            book_adjustments=Decimal("0"),
            period_from=date(2026, 4, 1),
            period_to=date(2026, 4, 30),
            stakeholder_rows=[self.first_row, self.second_row],
            balances={},
        )
        remuneration = next(line for line in result.lines if line.component_type == "remuneration")
        self.assertEqual(remuneration.amount, Decimal("1000.00"))

    def test_persisted_run_is_deterministic_and_idempotent(self):
        self.approve()
        kwargs = {
            "entity": self.entity,
            "entityfin": self.entityfin,
            "subentity": None,
            "period_from": date(2026, 4, 1),
            "period_to": date(2027, 3, 31),
            "cadence": "annual",
            "profit_source": "manual_approved",
            "supplied_profit": Decimal("100000"),
            "book_adjustments": Decimal("0"),
            "balance_inputs": [{"ownership": self.first.id, "capital_balance": "100000", "drawing_balance": "10000"}],
            "idempotency_key": "annual-2026-v1",
            "actor": self.maker,
        }
        first = calculate_distribution_run(**kwargs)
        repeated = calculate_distribution_run(**kwargs)
        self.assertEqual(first.id, repeated.id)
        self.assertEqual(first.status, CapitalDistributionRun.Status.CALCULATED)
        self.assertEqual(first.segments.count(), 1)
        self.assertEqual(first.lines.count(), 5)
        self.assertTrue(first.calculation_hash)

        changed = {**kwargs, "supplied_profit": Decimal("100001")}
        with self.assertRaises(ValidationError):
            calculate_distribution_run(**changed)

    def test_calculation_rejects_period_without_approved_policy_coverage(self):
        self.policy.effective_from = date(2026, 5, 1)
        self.policy.save(update_fields=("effective_from", "updated_at"))
        self.approve()
        with self.assertRaises(ValidationError):
            calculate_distribution_run(
                entity=self.entity,
                entityfin=self.entityfin,
                subentity=None,
                period_from=date(2026, 4, 1),
                period_to=date(2026, 5, 31),
                cadence="custom",
                profit_source="manual_approved",
                supplied_profit=Decimal("1000"),
                book_adjustments=Decimal("0"),
                balance_inputs=[],
                idempotency_key="gap-run",
                actor=self.maker,
            )

    def test_posted_books_source_is_captured_in_auditable_segment_snapshot(self):
        self.approve()
        with patch(
            "capital_distribution.services._segment_book_profit",
            return_value=(Decimal("2500.00"), {"source": "profit_loss_report", "net_profit": "2500.00"}),
        ):
            run = calculate_distribution_run(
                entity=self.entity,
                entityfin=self.entityfin,
                subentity=None,
                period_from=date(2026, 4, 1),
                period_to=date(2027, 3, 31),
                cadence="annual",
                profit_source="posted_books",
                supplied_profit=None,
                book_adjustments=Decimal("0"),
                balance_inputs=[],
                idempotency_key="posted-books-run",
                actor=self.maker,
            )
        self.assertEqual(run.source_profit, Decimal("2500.00"))
        self.assertEqual(run.segments.get().source_snapshot["source"], "profit_loss_report")


@override_settings(ROOT_URLCONF="FA.urls", AUTH_PASSWORD_VALIDATORS=[])
class CapitalDistributionRunAPITests(CapitalDistributionFixtureMixin, APITestCase):
    def setUp(self):
        self.maker = self.make_user("run-api-maker")
        self.approver = self.make_user("run-api-approver")
        self.entity, self.first, self.second = self.make_partnership(name="Run API Partnership", owner=self.maker)
        self.entityfin = EntityFinancialYear.objects.create(
            entity=self.entity,
            desc="FY 2026-27",
            year_code="FY2026-27",
            finstartyear=aware(2026, 4, 1),
            finendyear=aware(2027, 3, 31),
            createdby=self.maker,
        )
        self.profile = materialize_formation_profile(entity=self.entity, actor=self.maker)
        self.policy = create_policy(
            entity=self.entity,
            formation_profile=self.profile,
            payload=self.policy_payload(self.first, self.second, entityfin=self.entityfin),
            actor=self.maker,
        )
        first_row = self.policy.stakeholders.order_by("sort_order").first()
        first_row.configuration = {
            "remuneration": {"enabled": True, "method": "fixed", "amount": "12000", "prorate": True},
            "capital_interest": {"enabled": True, "rate": "10"},
            "drawing_interest": {"enabled": True, "rate": "12"},
        }
        first_row.save(update_fields=("configuration", "updated_at"))
        submit_policy(policy=self.policy, actor=self.maker)
        self.policy = approve_policy(policy=self.policy, actor=self.approver)
        self.client = APIClient()
        self.client.force_authenticate(self.maker)
        self.subscription_patch = patch("capital_distribution.views.SubscriptionService.assert_entity_access", return_value=None)
        self.permission_patch = patch(
            "capital_distribution.views.EffectivePermissionService.permission_codes_for_user",
            return_value={"capital_distribution.run.view", "capital_distribution.run.calculate"},
        )
        self.subscription_patch.start()
        self.permission_patch.start()
        self.addCleanup(self.subscription_patch.stop)
        self.addCleanup(self.permission_patch.stop)

    def test_calculate_list_and_detail_api(self):
        response = self.client.post(
            reverse("capital_distribution_api:run-list-calculate"),
            {
                "entity": self.entity.id,
                "entityfinid": self.entityfin.id,
                "period_from": "2026-04-01",
                "period_to": "2027-03-31",
                "cadence": "annual",
                "profit_source": "manual_approved",
                "source_profit": "100000.00",
                "book_adjustments": "0.00",
                "balances": [{"ownership": self.first.id, "capital_balance": "100000.00", "drawing_balance": "10000.00"}],
                "idempotency_key": "api-annual-2026-v1",
            },
            format="json",
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["distributable_result"], "79200.00")
        self.assertEqual(len(response.data["segments"]), 1)

        listed = self.client.get(reverse("capital_distribution_api:run-list-calculate"), {
            "entity": self.entity.id, "entityfinid": self.entityfin.id
        })
        self.assertEqual(listed.status_code, 200, listed.data)
        self.assertEqual(listed.data["count"], 1)

        detail = self.client.get(
            reverse("capital_distribution_api:run-detail", kwargs={"run_id": response.data["id"]}),
            {"entity": self.entity.id},
        )
        self.assertEqual(detail.status_code, 200, detail.data)
        self.assertEqual(detail.data["calculation_hash"], response.data["calculation_hash"])


class CapitalDistributionPostingLifecycleTests(CapitalDistributionFixtureMixin, TestCase):
    def setUp(self):
        self.maker = self.make_user("posting-maker")
        self.approver = self.make_user("posting-approver")
        self.entity, self.first, self.second = self.make_partnership(name="Posting Partnership", owner=self.maker)
        self.entityfin = EntityFinancialYear.objects.create(
            entity=self.entity,
            desc="FY 2026-27",
            year_code="FY2026-27",
            finstartyear=aware(2026, 4, 1),
            finendyear=aware(2027, 3, 31),
            createdby=self.maker,
        )
        self.profile = materialize_formation_profile(entity=self.entity, actor=self.maker)
        self.policy = create_policy(
            entity=self.entity,
            formation_profile=self.profile,
            payload=self.policy_payload(self.first, self.second, entityfin=self.entityfin),
            actor=self.maker,
        )
        submit_policy(policy=self.policy, actor=self.maker)
        self.policy = approve_policy(policy=self.policy, actor=self.approver)
        self.partner_accounts = {}
        for ownership in (self.first, self.second):
            self.partner_accounts[ownership.id] = self.make_account(f"{ownership.name} Current")
            upsert_account_mapping(
                entity=self.entity,
                ownership=ownership,
                capital_account=None,
                current_account=self.partner_accounts[ownership.id],
                drawings_account=None,
                effective_from=date(2026, 4, 1),
                effective_to=None,
                actor=self.maker,
            )
        self.appropriation = self.make_account("Profit and Loss Appropriation")
        static, _ = StaticAccount.objects.update_or_create(
            code="PROFIT_LOSS_APPROPRIATION",
            defaults={"name": "Profit and Loss Appropriation", "group": StaticAccountGroup.EQUITY},
        )
        EntityStaticAccountMap.objects.create(
            entity=self.entity,
            static_account=static,
            account=self.appropriation,
            ledger=self.appropriation.ledger,
            createdby=self.maker,
        )

    def make_account(self, name, *, opening_credit="0.00", opening_debit="0.00"):
        ledger = Ledger.objects.create(
            entity=self.entity,
            name=name,
            openingbcr=opening_credit,
            openingbdr=opening_debit,
            createdby=self.maker,
        )
        return account.objects.create(entity=self.entity, accountname=name, ledger=ledger, createdby=self.maker)

    def calculate(
        self,
        *,
        key="posting-run",
        balances=None,
        subentity=None,
        period_from=date(2026, 4, 1),
        period_to=date(2027, 3, 31),
        supplied_profit=Decimal("100000.00"),
        cadence="annual",
    ):
        if balances is None:
            balances = [
                {"ownership": self.first.id, "capital_balance": "100000.00"},
                {"ownership": self.second.id, "capital_balance": "50000.00"},
            ]
        return calculate_distribution_run(
            entity=self.entity,
            entityfin=self.entityfin,
            subentity=subentity,
            period_from=period_from,
            period_to=period_to,
            cadence=cadence,
            profit_source="manual_approved",
            supplied_profit=supplied_profit,
            book_adjustments=Decimal("0.00"),
            balance_inputs=balances,
            idempotency_key=key,
            actor=self.maker,
        )

    def post_run(self, **kwargs):
        run = self.calculate(**kwargs)
        run = submit_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
        run = approve_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        return post_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)

    def test_posted_appropriation_closes_and_carries_partner_balances_once(self):
        run = self.post_run(key="year-opening-certification")
        close_snapshot = {
            "financial_year": self.entityfin,
            "pnl": {
                "income": [{
                    "label": "Certified operating profit",
                    "accounthead_id": 990001,
                    "debit": "0.00",
                    "credit": "100000.00",
                    "amount_decimal": "100000.00",
                }],
                "expenses": [],
            },
            "summary": {"net_profit": Decimal("100000.00")},
        }
        adapter_context = {
            "validation_issues": [],
            "equity_targets": [],
            "missing_equity_codes": [],
            "equity_allocation_mode": "ratio_split",
            "constitution": {"constitution_mode": "partnership"},
            "allocation_plan": [],
        }

        with patch(
            "reports.services.controls.year_end_close.YearOpeningPostingAdapter.build_context",
            return_value=adapter_context,
        ):
            close_lines, close_meta, diagnostics = _build_close_journal_lines(
                snapshot=close_snapshot,
                entity_id=self.entity.id,
                entityfin_id=self.entityfin.id,
                subentity_id=None,
                opening_policy={},
            )

        appropriation_close_lines = [
            line for line in close_lines if line.ledger_id == self.appropriation.ledger_id
        ]
        self.assertEqual(len(appropriation_close_lines), 1)
        self.assertFalse(appropriation_close_lines[0].drcr)
        self.assertEqual(appropriation_close_lines[0].amount, Decimal("100000.00"))
        self.assertEqual(diagnostics["equity_allocation_mode"], "posted_appropriation_clearance")
        self.assertEqual(
            [row["source"] for row in close_meta if row["section"] == "equity"],
            ["equity_target"],
        )

        posted_appropriation_debit = sum(
            JournalLine.objects.filter(
                entity=self.entity,
                txn_type=TxnType.CAPITAL_DISTRIBUTION,
                txn_id=run.id,
                ledger=self.appropriation.ledger,
                drcr=True,
            ).values_list("amount", flat=True),
            Decimal("0.00"),
        )
        self.assertEqual(posted_appropriation_debit, appropriation_close_lines[0].amount)

        opening_snapshot = {
            "bs": {
                "assets": [{
                    "ledger_id": 880001,
                    "accounthead_id": 880001,
                    "ledger_name": "Certified cash balance",
                    "amount_decimal": "100000.00",
                }],
                "liabilities_and_equity": [
                    {
                        "ledger_id": self.partner_accounts[self.first.id].ledger_id,
                        "accounthead_id": 880002,
                        "ledger_name": self.partner_accounts[self.first.id].ledger.name,
                        "amount_decimal": "60000.00",
                    },
                    {
                        "ledger_id": self.partner_accounts[self.second.id].ledger_id,
                        "accounthead_id": 880002,
                        "ledger_name": self.partner_accounts[self.second.id].ledger.name,
                        "amount_decimal": "40000.00",
                    },
                ],
                "summary": {
                    "net_profit_brought_to_equity": "0.00",
                    "raw_net_profit": "0.00",
                },
                "stock_valuation": {},
            }
        }
        opening_context = {
            "destination_ledgers": {
                "equity": {"static_account_code": "OPENING_EQUITY_TRANSFER", "ledger_id": None},
                "inventory": {"static_account_code": "OPENING_INVENTORY_CARRY_FORWARD", "ledger_id": 880003},
            },
            "constitution": {"constitution_mode": "partnership"},
            "allocation_plan": [],
            "equity_targets": [],
            "missing_equity_codes": [],
            "equity_allocation_mode": "ratio_split",
            "validation_issues": [],
        }
        with patch(
            "reports.services.controls.opening_generation.YearOpeningPostingAdapter.build_context",
            return_value=opening_context,
        ):
            opening_lines, opening_meta, opening_summary = _build_opening_lines(
                opening_snapshot,
                opening_policy={},
                entity_id=self.entity.id,
            )

        partner_credits = {
            line.ledger_id: line.amount
            for line in opening_lines
            if not line.drcr and line.ledger_id in {
                self.partner_accounts[self.first.id].ledger_id,
                self.partner_accounts[self.second.id].ledger_id,
            }
        }
        self.assertEqual(partner_credits, {
            self.partner_accounts[self.first.id].ledger_id: Decimal("60000.00"),
            self.partner_accounts[self.second.id].ledger_id: Decimal("40000.00"),
        })
        self.assertEqual(sum(line.amount for line in opening_lines if line.drcr), Decimal("100000.00"))
        self.assertEqual(sum(line.amount for line in opening_lines if not line.drcr), Decimal("100000.00"))
        self.assertFalse(any(row["source"].startswith("synthetic_") for row in opening_meta))
        self.assertEqual(opening_summary["diagnostics"]["synthetic_equity_adjustment"], "0.00")

    def test_year_end_coverage_rejects_partial_and_ignores_reversed_runs(self):
        partial = self.post_run(
            key="partial-year-certification",
            period_from=date(2026, 4, 1),
            period_to=date(2026, 9, 30),
            supplied_profit=Decimal("50000.00"),
            cadence="custom",
        )
        with self.assertRaisesMessage(DRFValidationError, "do not cover the full financial year"):
            _posted_appropriation_coverage(
                entity_id=self.entity.id,
                entityfin_id=self.entityfin.id,
                subentity_id=None,
                period_from=date(2026, 4, 1),
                period_to=date(2027, 3, 31),
            )

        partial = reverse_distribution_run(
            run=partial,
            actor=self.approver,
            expected_updated_at=partial.updated_at,
            reason="Certification reversal",
        )
        self.assertEqual(partial.status, CapitalDistributionRun.Status.REVERSED)
        self.assertIsNone(_posted_appropriation_coverage(
            entity_id=self.entity.id,
            entityfin_id=self.entityfin.id,
            subentity_id=None,
            period_from=date(2026, 4, 1),
            period_to=date(2027, 3, 31),
        ))

    def test_year_end_close_rejects_posted_appropriation_amount_mismatch(self):
        self.post_run(
            key="mismatched-year-certification",
            supplied_profit=Decimal("90000.00"),
        )
        snapshot = {
            "financial_year": self.entityfin,
            "pnl": {
                "income": [{
                    "label": "Certified operating profit",
                    "accounthead_id": 990001,
                    "debit": "0.00",
                    "credit": "100000.00",
                    "amount_decimal": "100000.00",
                }],
                "expenses": [],
            },
            "summary": {"net_profit": Decimal("100000.00")},
        }
        adapter_context = {
            "validation_issues": [],
            "equity_targets": [],
            "missing_equity_codes": [],
            "equity_allocation_mode": "ratio_split",
            "constitution": {"constitution_mode": "partnership"},
            "allocation_plan": [],
        }

        with (
            patch(
                "reports.services.controls.year_end_close.YearOpeningPostingAdapter.build_context",
                return_value=adapter_context,
            ),
            self.assertRaisesMessage(
                DRFValidationError,
                "Posted capital-distribution total does not match the closeable book result",
            ),
        ):
            _build_close_journal_lines(
                snapshot=snapshot,
                entity_id=self.entity.id,
                entityfin_id=self.entityfin.id,
                subentity_id=None,
                opening_policy={},
            )

    def test_calculate_submit_approve_post_retry_and_exact_reverse(self):
        run = self.calculate()
        run = submit_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
        with self.assertRaises(ValidationError):
            approve_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
        run = approve_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        run = post_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        original_batch_id = run.posting_batch_id
        posted = list(JournalLine.objects.filter(
            entity=self.entity,
            txn_type=TxnType.CAPITAL_DISTRIBUTION,
            txn_id=run.id,
        ).values_list("ledger_id", "drcr", "amount", "detail_id"))
        self.assertTrue(posted)
        self.assertEqual(
            sum(row[2] for row in posted if row[1]),
            sum(row[2] for row in posted if not row[1]),
        )
        coverage = _posted_appropriation_coverage(
            entity_id=self.entity.id,
            entityfin_id=self.entityfin.id,
            subentity_id=None,
            period_from=date(2026, 4, 1),
            period_to=date(2027, 3, 31),
        )
        self.assertEqual(coverage["net_allocation"], Decimal("100000.00"))

        retried = post_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        self.assertEqual(retried.posting_batch_id, original_batch_id)

        run = reverse_distribution_run(
            run=retried,
            actor=self.approver,
            expected_updated_at=retried.updated_at,
            reason="Correction required",
        )
        reversed_rows = list(JournalLine.objects.filter(
            entity=self.entity,
            txn_type=TxnType.CAPITAL_DISTRIBUTION,
            txn_id=run.id,
        ).values_list("ledger_id", "drcr", "amount", "detail_id"))
        self.assertEqual(
            sorted((ledger, not side, amount, detail) for ledger, side, amount, detail in posted),
            sorted(reversed_rows),
        )
        self.assertNotEqual(run.reversal_batch_id, original_batch_id)
        self.assertEqual(run.status, CapitalDistributionRun.Status.REVERSED)

    def test_appropriation_statement_reconciles_posting_and_exact_reversal(self):
        run = self.calculate(key="statement-run")
        draft_statement = build_appropriation_statement(
            entity=self.entity,
            entityfin=self.entityfin,
            period_from=date(2026, 4, 1),
            period_to=date(2027, 3, 31),
        )
        self.assertEqual(draft_statement["runs"][0]["reconciliation_status"], "not_posted")
        self.assertIsNone(draft_statement["runs"][0]["reconciled"])

        run = submit_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
        run = approve_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        run = post_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        posted_statement = build_appropriation_statement(
            entity=self.entity,
            entityfin=self.entityfin,
            period_from=date(2026, 4, 1),
            period_to=date(2027, 3, 31),
        )
        posted_run = posted_statement["runs"][0]
        self.assertEqual(posted_run["reconciliation_status"], "reconciled")
        self.assertEqual(posted_run["calculated_net"], "100000.00")
        self.assertEqual(posted_run["effective_net"], "100000.00")
        self.assertEqual(posted_run["original_journal"]["debit"], "100000.00")
        self.assertEqual(posted_run["original_journal"]["credit"], "100000.00")
        self.assertEqual(posted_statement["summary"]["effective_credit"], "100000.00")
        self.assertEqual(
            {row["stakeholder"]: row["effective_credit"] for row in posted_statement["partners"]},
            {self.first.name: "60000.00", self.second.name: "40000.00"},
        )

        run = reverse_distribution_run(
            run=run,
            actor=self.approver,
            expected_updated_at=run.updated_at,
            reason="Statement reversal test",
        )
        reversed_statement = build_appropriation_statement(
            entity=self.entity,
            entityfin=self.entityfin,
            period_from=date(2026, 4, 1),
            period_to=date(2027, 3, 31),
        )
        reversed_run = reversed_statement["runs"][0]
        self.assertEqual(reversed_run["reconciliation_status"], "reconciled", reversed_run)
        self.assertEqual(reversed_run["effective_net"], "0.00")
        self.assertFalse(reversed_run["original_journal"]["available"])
        self.assertTrue(reversed_run["original_journal"]["reconstructed_from_frozen_run"])
        self.assertEqual(reversed_statement["summary"]["effective_credit"], "0.00")
        self.assertTrue(all(row["effective_debit"] == "0.00" for row in reversed_statement["partners"]))
        self.assertTrue(all(row["effective_credit"] == "0.00" for row in reversed_statement["partners"]))

    def test_financial_reports_use_posted_journals_not_draft_calculations(self):
        from reports.services.financial.ledger_book import build_ledger_book
        from reports.services.financial.statements import build_balance_sheet, build_profit_and_loss
        from reports.services.financial.trial_balance import build_trial_balance

        run = self.calculate(key="financial-report-boundary")
        report_args = {
            "entity_id": self.entity.id,
            "entityfin_id": self.entityfin.id,
            "from_date": "2026-04-01",
            "to_date": "2027-03-31",
        }
        draft_profit_loss = build_profit_and_loss(**report_args)
        draft_disclosure = draft_profit_loss["capital_distribution"]
        self.assertEqual(draft_disclosure["posted_accounting"]["run_count"], 0)
        self.assertEqual(draft_disclosure["posted_accounting"]["journal_debit"], "0.00")
        self.assertEqual(draft_disclosure["pending_calculation"]["run_count"], 1)
        self.assertFalse(draft_disclosure["pending_calculation"]["affects_books"])

        run = submit_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
        run = approve_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        approved_profit_loss = build_profit_and_loss(**report_args)
        self.assertEqual(approved_profit_loss["totals"]["net_profit"], draft_profit_loss["totals"]["net_profit"])
        self.assertEqual(approved_profit_loss["capital_distribution"]["posted_accounting"]["run_count"], 0)

        run = post_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        posted_profit_loss = build_profit_and_loss(**report_args)
        posted_disclosure = posted_profit_loss["capital_distribution"]
        self.assertEqual(posted_profit_loss["totals"]["net_profit"], draft_profit_loss["totals"]["net_profit"])
        self.assertEqual(posted_disclosure["operating_profit_impact"], "0.00")
        self.assertEqual(posted_disclosure["posted_accounting"]["run_count"], 1)
        self.assertEqual(posted_disclosure["posted_accounting"]["journal_debit"], "100000.00")
        self.assertEqual(posted_disclosure["posted_accounting"]["journal_credit"], "100000.00")
        self.assertEqual(posted_disclosure["posted_accounting"]["net_equity_movement"], "100000.00")
        self.assertTrue(posted_disclosure["posted_accounting"]["balanced"])

        balance_sheet = build_balance_sheet(**report_args)
        self.assertEqual(
            balance_sheet["capital_distribution"]["posted_accounting"],
            posted_disclosure["posted_accounting"],
        )
        from reports.api.financial.views import _balance_sheet_export_meta, _profit_loss_export_meta

        scope_names = {
            "entity_name": self.entity.entityname,
            "entityfin_name": self.entityfin.desc,
            "subentity_name": None,
        }
        export_scope = {
            "scope_mode": "financial_year",
            "posted_only": True,
            "hide_zero_rows": True,
        }
        self.assertIn(
            "Posted Appropriation",
            dict(_profit_loss_export_meta(scope_names, export_scope, posted_profit_loss)),
        )
        self.assertIn(
            "Partner Equity Movement",
            dict(_balance_sheet_export_meta(scope_names, export_scope, balance_sheet)),
        )

        ledger_ids = [self.appropriation.ledger_id] + [
            row.ledger_id for row in self.partner_accounts.values()
        ]
        trial_balance = build_trial_balance(**report_args, ledger_ids=ledger_ids)
        self.assertEqual(trial_balance["totals"]["debit"], "100000.00")
        self.assertEqual(trial_balance["totals"]["credit"], "100000.00")

        first_partner_book = build_ledger_book(
            **report_args,
            ledger_id=self.partner_accounts[self.first.id].ledger_id,
        )
        self.assertEqual(first_partner_book["totals"]["debit"], "0.00")
        self.assertEqual(first_partner_book["totals"]["credit"], "60000.00")
        self.assertEqual(first_partner_book["rows"][0]["voucher_type"], TxnType.CAPITAL_DISTRIBUTION)

        run = reverse_distribution_run(
            run=run,
            actor=self.approver,
            expected_updated_at=run.updated_at,
            reason="Financial report boundary reversal",
        )
        reversed_profit_loss = build_profit_and_loss(**report_args)
        reversed_posting = reversed_profit_loss["capital_distribution"]["posted_accounting"]
        self.assertEqual(reversed_profit_loss["totals"]["net_profit"], draft_profit_loss["totals"]["net_profit"])
        self.assertEqual(reversed_posting["run_count"], 0)
        self.assertEqual(reversed_posting["journal_debit"], "0.00")
        self.assertEqual(reversed_posting["net_equity_movement"], "0.00")

    @override_settings(ROOT_URLCONF="FA.urls")
    def test_financial_report_exports_render_posted_appropriation_metadata(self):
        from io import BytesIO

        from openpyxl import load_workbook
        from pypdf import PdfReader

        run = self.calculate(key="financial-export-certification")
        run = submit_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
        run = approve_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
        post_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)

        client = APIClient()
        client.force_authenticate(self.approver)
        params = {
            "entity": self.entity.id,
            "entityfinid": self.entityfin.id,
            "from_date": "2026-04-01",
            "to_date": "2027-03-31",
        }
        report_cases = (
            ("profit-loss", "Posted Appropriation"),
            ("balance-sheet", "Partner Equity Movement"),
        )
        permissions = {
            "reports.financial_hub.profit_loss.view",
            "reports.financial_hub.balance_sheet.view",
        }
        with (
            patch("core.entitlements.SubscriptionService.assert_entity_access", return_value=None),
            patch("core.entitlements.EffectivePermissionService.has_scope_access", return_value=True),
            patch("core.entitlements.EffectivePermissionService.has_data_scope_access", return_value=True),
            patch(
                "reports.api.report_permissions.EffectivePermissionService.permission_codes_for_user",
                return_value=permissions,
            ),
        ):
            for report_name, movement_label in report_cases:
                csv_response = client.get(
                    reverse(f"reports_api:financial-{report_name}-csv"),
                    params,
                )
                self.assertEqual(csv_response.status_code, 200, csv_response.content[:500])
                csv_text = csv_response.content.decode("utf-8-sig")
                self.assertIn(movement_label, csv_text)
                self.assertIn("Appropriation Reconciliation", csv_text)
                self.assertIn("Balanced", csv_text)

                excel_response = client.get(
                    reverse(f"reports_api:financial-{report_name}-excel"),
                    params,
                )
                self.assertEqual(excel_response.status_code, 200, excel_response.content[:500])
                workbook = load_workbook(BytesIO(excel_response.content), data_only=True)
                excel_text = "\n".join(
                    str(cell.value)
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.value is not None
                )
                self.assertIn(movement_label, excel_text)
                self.assertIn("Appropriation Reconciliation", excel_text)
                self.assertIn("Balanced", excel_text)

                pdf_response = client.get(
                    reverse(f"reports_api:financial-{report_name}-pdf"),
                    params,
                )
                self.assertEqual(pdf_response.status_code, 200, pdf_response.content[:500])
                self.assertTrue(pdf_response.content.startswith(b"%PDF"))
                reader = PdfReader(BytesIO(pdf_response.content))
                pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
                self.assertIn(movement_label, pdf_text)
                self.assertIn("Appropriation Reconciliation", pdf_text)
                self.assertIn("Balanced", pdf_text)

                print_response = client.get(
                    reverse(f"reports_api:financial-{report_name}-print"),
                    params,
                )
                self.assertEqual(print_response.status_code, 200, print_response.content[:500])
                self.assertTrue(print_response["Content-Disposition"].startswith("inline;"))
                self.assertTrue(print_response.content.startswith(b"%PDF"))

    def test_financial_disclosure_isolates_branches_and_consolidates_posted_runs(self):
        from capital_distribution.reporting import build_financial_statement_disclosure

        north = SubEntity.objects.create(
            entity=self.entity,
            subentityname="North Branch",
            subentity_code="NORTH",
            branch_type=SubEntity.BranchType.BRANCH,
        )
        south = SubEntity.objects.create(
            entity=self.entity,
            subentityname="South Branch",
            subentity_code="SOUTH",
            branch_type=SubEntity.BranchType.BRANCH,
        )

        def post_branch(branch, amount, key):
            run = self.calculate(
                key=key,
                subentity=branch,
                supplied_profit=Decimal(amount),
            )
            run = submit_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
            run = approve_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
            return post_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)

        north_run = post_branch(north, "1000.00", "north-branch-run")
        post_branch(south, "2000.00", "south-branch-run")
        disclosure_args = {
            "entity_id": self.entity.id,
            "entityfin_id": self.entityfin.id,
            "period_from": date(2026, 4, 1),
            "period_to": date(2027, 3, 31),
        }

        north_report = build_financial_statement_disclosure(**disclosure_args, subentity_id=north.id)
        south_report = build_financial_statement_disclosure(**disclosure_args, subentity_id=south.id)
        consolidated = build_financial_statement_disclosure(**disclosure_args)
        self.assertEqual(north_report["posted_accounting"]["run_count"], 1)
        self.assertEqual(north_report["posted_accounting"]["net_equity_movement"], "1000.00")
        self.assertEqual(south_report["posted_accounting"]["run_count"], 1)
        self.assertEqual(south_report["posted_accounting"]["net_equity_movement"], "2000.00")
        self.assertEqual(consolidated["posted_accounting"]["run_count"], 2)
        self.assertEqual(consolidated["posted_accounting"]["net_equity_movement"], "3000.00")
        self.assertTrue(consolidated["posted_accounting"]["balanced"])

        reverse_distribution_run(
            run=north_run,
            actor=self.approver,
            expected_updated_at=north_run.updated_at,
            reason="Branch consolidation certification",
        )
        reversed_north = build_financial_statement_disclosure(**disclosure_args, subentity_id=north.id)
        remaining = build_financial_statement_disclosure(**disclosure_args)
        self.assertEqual(reversed_north["posted_accounting"]["run_count"], 0)
        self.assertEqual(reversed_north["posted_accounting"]["net_equity_movement"], "0.00")
        self.assertEqual(remaining["posted_accounting"]["run_count"], 1)
        self.assertEqual(remaining["posted_accounting"]["net_equity_movement"], "2000.00")

    def test_financial_disclosure_filters_multiple_periodic_postings_by_date(self):
        from capital_distribution.reporting import build_financial_statement_disclosure

        def post_period(period_from, period_to, amount, key):
            run = self.calculate(
                key=key,
                period_from=period_from,
                period_to=period_to,
                supplied_profit=Decimal(amount),
                cadence="monthly",
            )
            run = submit_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
            run = approve_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)
            return post_distribution_run(run=run, actor=self.approver, expected_updated_at=run.updated_at)

        post_period(date(2026, 4, 1), date(2026, 4, 30), "1000.00", "april-periodic-run")
        post_period(date(2026, 5, 1), date(2026, 5, 31), "500.00", "may-periodic-run")
        common = {"entity_id": self.entity.id, "entityfin_id": self.entityfin.id}

        april = build_financial_statement_disclosure(
            **common,
            period_from=date(2026, 4, 1),
            period_to=date(2026, 4, 30),
        )
        year_to_date = build_financial_statement_disclosure(
            **common,
            period_from=date(2026, 4, 1),
            period_to=date(2026, 5, 31),
        )
        self.assertEqual(april["posted_accounting"]["run_count"], 1)
        self.assertEqual(april["posted_accounting"]["net_equity_movement"], "1000.00")
        self.assertEqual(april["audit_reconciliation"]["reconciled_run_count"], 1)
        self.assertEqual(year_to_date["posted_accounting"]["run_count"], 2)
        self.assertEqual(year_to_date["posted_accounting"]["net_equity_movement"], "1500.00")
        self.assertEqual(year_to_date["audit_reconciliation"]["reconciled_run_count"], 2)

    @override_settings(ROOT_URLCONF="FA.urls")
    def test_appropriation_statement_api_enforces_scope_and_date_validation(self):
        self.calculate(key="statement-api-run")
        client = APIClient()
        client.force_authenticate(self.maker)
        with (
            patch("capital_distribution.views.SubscriptionService.assert_entity_access", return_value=None),
            patch(
                "capital_distribution.views.EffectivePermissionService.permission_codes_for_user",
                return_value={"capital_distribution.run.view"},
            ),
        ):
            response = client.get(reverse("capital_distribution_api:appropriation-statement"), {
                "entity": self.entity.id,
                "entityfinid": self.entityfin.id,
                "period_from": "2026-04-01",
                "period_to": "2027-03-31",
            })
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data["summary"]["run_count"], 1)

            invalid = client.get(reverse("capital_distribution_api:appropriation-statement"), {
                "entity": self.entity.id,
                "entityfinid": self.entityfin.id,
                "period_from": "2027-03-31",
                "period_to": "2026-04-01",
            })
            self.assertEqual(invalid.status_code, 400, invalid.data)

    def test_posted_ledger_balance_change_blocks_submission_as_stale(self):
        self.partner_accounts[self.first.id].ledger.openingbcr = Decimal("100000.00")
        self.partner_accounts[self.first.id].ledger.save(update_fields=("openingbcr",))
        run = self.calculate(key="auto-balance-run", balances=[])
        self.partner_accounts[self.first.id].ledger.openingbcr = Decimal("100001.00")
        self.partner_accounts[self.first.id].ledger.save(update_fields=("openingbcr",))
        with self.assertRaises(ValidationError) as captured:
            submit_distribution_run(run=run, actor=self.maker, expected_updated_at=run.updated_at)
        self.assertIn("source", captured.exception.message_dict)

    def test_mapping_rejects_cross_entity_account(self):
        other = self.make_entity(name="Other Mapping Entity", owner=self.maker)
        ledger = Ledger.objects.create(entity=other, name="Wrong Current", createdby=self.maker)
        wrong = account.objects.create(entity=other, accountname="Wrong Current", ledger=ledger, createdby=self.maker)
        with self.assertRaises(ValidationError):
            upsert_account_mapping(
                entity=self.entity,
                ownership=self.first,
                capital_account=None,
                current_account=wrong,
                drawings_account=None,
                effective_from=None,
                effective_to=None,
                actor=self.maker,
            )

    @override_settings(ROOT_URLCONF="FA.urls")
    def test_run_direct_objects_enforce_the_persisted_branch_scope(self):
        allowed_branch = SubEntity.objects.create(entity=self.entity, subentityname="Allowed Branch")
        foreign_branch = SubEntity.objects.create(entity=self.entity, subentityname="Foreign Branch")
        allowed_run = self.calculate(key="allowed-branch-run", subentity=allowed_branch)
        foreign_run = self.calculate(key="foreign-branch-run", subentity=foreign_branch)
        client = APIClient()
        client.force_authenticate(self.maker)

        with (
            patch("capital_distribution.views.SubscriptionService.assert_entity_access", return_value=None),
            patch(
                "capital_distribution.views.EffectivePermissionService.permission_codes_for_user",
                return_value={"capital_distribution.run.view", "capital_distribution.run.submit"},
            ),
            patch(
                "capital_distribution.views.EffectivePermissionService.has_scope_access",
                side_effect=lambda _user, _entity_id, subentity_id=None: subentity_id == allowed_branch.id,
            ),
        ):
            listed = client.get(
                reverse("capital_distribution_api:run-list-calculate"),
                {"entity": self.entity.id, "entityfinid": self.entityfin.id, "subentity": allowed_branch.id},
            )
            self.assertEqual(listed.status_code, 200, listed.data)
            self.assertEqual([row["id"] for row in listed.data["results"]], [allowed_run.id])

            detail = client.get(
                reverse("capital_distribution_api:run-detail", kwargs={"run_id": allowed_run.id}),
                {"entity": self.entity.id},
            )
            self.assertEqual(detail.status_code, 200, detail.data)

            submitted = client.post(
                reverse("capital_distribution_api:run-submit", kwargs={"run_id": allowed_run.id}),
                {"entity": self.entity.id, "expected_updated_at": allowed_run.updated_at.isoformat()},
                format="json",
            )
            self.assertEqual(submitted.status_code, 200, submitted.data)

            foreign_detail = client.get(
                reverse("capital_distribution_api:run-detail", kwargs={"run_id": foreign_run.id}),
                {"entity": self.entity.id},
            )
            self.assertEqual(foreign_detail.status_code, 403, foreign_detail.data)

            aggregate = client.get(
                reverse("capital_distribution_api:run-list-calculate"),
                {"entity": self.entity.id, "entityfinid": self.entityfin.id},
            )
            self.assertEqual(aggregate.status_code, 403, aggregate.data)

    @override_settings(ROOT_URLCONF="FA.urls")
    def test_scoped_api_completes_run_lifecycle_and_requires_separate_approver(self):
        run = self.calculate(key="api-posting-run")
        client = APIClient()
        permissions = {
            "capital_distribution.run.submit",
            "capital_distribution.run.approve",
            "capital_distribution.run.post",
            "capital_distribution.run.reverse",
        }
        with (
            patch("capital_distribution.views.SubscriptionService.assert_entity_access", return_value=None),
            patch(
                "capital_distribution.views.EffectivePermissionService.permission_codes_for_user",
                return_value=permissions,
            ),
        ):
            client.force_authenticate(self.maker)
            submitted = client.post(
                reverse("capital_distribution_api:run-submit", kwargs={"run_id": run.id}),
                {"entity": self.entity.id, "expected_updated_at": run.updated_at.isoformat()},
                format="json",
            )
            self.assertEqual(submitted.status_code, 200, submitted.data)
            self.assertEqual(submitted.data["status"], CapitalDistributionRun.Status.SUBMITTED)

            self_approval = client.post(
                reverse("capital_distribution_api:run-approve", kwargs={"run_id": run.id}),
                {"entity": self.entity.id, "expected_updated_at": submitted.data["updated_at"]},
                format="json",
            )
            self.assertEqual(self_approval.status_code, 400, self_approval.data)

            client.force_authenticate(self.approver)
            approved = client.post(
                reverse("capital_distribution_api:run-approve", kwargs={"run_id": run.id}),
                {"entity": self.entity.id, "expected_updated_at": submitted.data["updated_at"]},
                format="json",
            )
            self.assertEqual(approved.status_code, 200, approved.data)
            posted = client.post(
                reverse("capital_distribution_api:run-post", kwargs={"run_id": run.id}),
                {"entity": self.entity.id, "expected_updated_at": approved.data["updated_at"]},
                format="json",
            )
            self.assertEqual(posted.status_code, 200, posted.data)
            self.assertEqual(posted.data["status"], CapitalDistributionRun.Status.POSTED)
            self.assertIsNotNone(posted.data["posting_batch"])

            reversed_run = client.post(
                reverse("capital_distribution_api:run-reverse", kwargs={"run_id": run.id}),
                {
                    "entity": self.entity.id,
                    "expected_updated_at": posted.data["updated_at"],
                    "reason": "API correction test",
                },
                format="json",
            )
            self.assertEqual(reversed_run.status_code, 200, reversed_run.data)
            self.assertEqual(reversed_run.data["status"], CapitalDistributionRun.Status.REVERSED)
            self.assertIsNotNone(reversed_run.data["reversal_batch"])
